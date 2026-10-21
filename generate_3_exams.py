#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generate_3_exams.py
Tạo 3 đề kiểm tra định kỳ giữa HK1 môn Khoa học tự nhiên (Lớp 7, 8, 9)
theo đúng chuẩn FORM ĐỀ MẪU CHUẨN (A4, TNR 12pt, lề 2.5/2.0cm, 7.5cm tab stop,
phân bố khóa trắc nghiệm 4A-4B-4C-4D, bảng Đúng/Sai, thang điểm tự luận 4 ý x 0.25đ).
"""

import os
import sys
from PIL import ImageFont
import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_NAME = "Times New Roman"
FONT_SIZE_PT = 12

COLOR_BLACK = RGBColor(0, 0, 0)
COLOR_NAVY = RGBColor(0x00, 0x33, 0x66)       # 003366 tiêu đề (Xanh đậm)
COLOR_RED_LABEL = RGBColor(0xB4, 0x00, 0x00) # B40000 đáp án trắc nghiệm (Đỏ sẫm)
COLOR_RED_SAI = RGBColor(0xC8, 0x00, 0x00)   # C80000 chữ "SAI"
COLOR_GREEN = RGBColor(0x00, 0x80, 0x00)     # 008000 chữ "ĐÚNG"
COLOR_GRAY = RGBColor(0x78, 0x78, 0x78)      # 787878 đường nét đứt

HEX_BG_GRAY = "EAECEE"
HEX_BG_BLUE_LIGHT = "D6EAF8"
HEX_BG_HEADER_LIGHT = "F2F4F4"
HEX_BORDER = "000000"  # Viền đen theo yêu cầu PHẦN B

LIBERATION_SERIF_PATH = "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"
if not os.path.exists(LIBERATION_SERIF_PATH):
    LIBERATION_SERIF_PATH = None


def measure_text_cm(text: str) -> float:
    """Đo độ rộng text bằng Pillow ở phông 12pt (16px tại 96dpi)"""
    if not text:
        return 0.0
    if LIBERATION_SERIF_PATH:
        try:
            font = ImageFont.truetype(LIBERATION_SERIF_PATH, 16)
            bbox = font.getbbox(text)
            w_px = bbox[2] - bbox[0]
            return (w_px * 2.54) / 96.0
        except Exception:
            pass
    return len(text) * 0.20


def format_run(run, bold=False, italic=False, color=COLOR_BLACK, size_pt=FONT_SIZE_PT, subscript=False, superscript=False):
    """Định dạng run chuẩn Times New Roman cho cả 4 thuộc tính (ascii, hAnsi, eastAsia, cs)"""
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    rPr = run._element.get_or_add_rPr()
    rFonts = parse_xml(
        r'<w:rFonts %s w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="Times New Roman" w:cs="Times New Roman"/>' % nsdecls('w')
    )
    rPr.append(rFonts)


def set_cell_shading(cell, color_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top_pt=3, bottom_pt=3, left_pt=5, right_pt=5):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{int(top_pt * 20)}" w:type="dxa"/>'
        f'<w:bottom w:w="{int(bottom_pt * 20)}" w:type="dxa"/>'
        f'<w:left w:w="{int(left_pt * 20)}" w:type="dxa"/>'
        f'<w:right w:w="{int(right_pt * 20)}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, border_color=HEX_BORDER, border_size="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f"""
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="{border_size}" w:space="0" w:color="{border_color}"/>
                <w:bottom w:val="single" w:sz="{border_size}" w:space="0" w:color="{border_color}"/>
                <w:left w:val="single" w:sz="{border_size}" w:space="0" w:color="{border_color}"/>
                <w:right w:val="single" w:sz="{border_size}" w:space="0" w:color="{border_color}"/>
                <w:insideH w:val="single" w:sz="{border_size}" w:space="0" w:color="{border_color}"/>
                <w:insideV w:val="single" w:sz="{border_size}" w:space="0" w:color="{border_color}"/>
            </w:tblBorders>
        """)
        tblPr[0].append(borders)


def set_row_cant_split(row):
    trPr = row._element.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def set_row_header(row):
    trPr = row._element.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))


def create_document():
    doc = docx.Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    style = doc.styles['Normal']
    style.font.name = FONT_NAME
    style.font.size = Pt(FONT_SIZE_PT)
    style.font.color.rgb = COLOR_BLACK
    rFonts = parse_xml(
        r'<w:rFonts %s w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="Times New Roman" w:cs="Times New Roman"/>' % nsdecls('w')
    )
    style.element.rPr.append(rFonts)
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(2)

    return doc


def add_exam_header(doc, school_info, exam_info):
    """
    PHẦN C.1: Bảng 2 cột không viền -> Đường kẻ nét đứt xám -> Dòng học sinh -> Tiêu đề
    """
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c0 = tbl.cell(0, 0)
    c1 = tbl.cell(0, 1)
    c0.width = Cm(7.2)
    c1.width = Cm(9.3)
    set_cell_margins(c0, 0, 2, 0, 5)
    set_cell_margins(c1, 0, 2, 5, 0)

    # Cột trái
    p_left = c0.paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_left.paragraph_format.line_spacing = 1.0
    p_left.paragraph_format.space_after = Pt(2)
    r1 = p_left.add_run(school_info['dept'] + "\n")
    format_run(r1, bold=True)
    r2 = p_left.add_run(school_info['school'] + "\n")
    format_run(r2, bold=True)
    r3 = p_left.add_run(f"MÃ ĐỀ: {exam_info['code']}")
    format_run(r3, bold=True)

    # Cột phải
    p_right = c1.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_right.paragraph_format.line_spacing = 1.0
    p_right.paragraph_format.space_after = Pt(2)
    r4 = p_right.add_run(exam_info['title'] + "\n")
    format_run(r4, bold=True)
    r5 = p_right.add_run(f"NĂM HỌC {exam_info['year']}\n")
    format_run(r5, bold=True)
    r6 = p_right.add_run(f"Môn: {exam_info['subject']} - LỚP {exam_info['grade']}\n")
    format_run(r6, bold=True)
    r7 = p_right.add_run(f"Thời gian làm bài: {exam_info['duration']} phút (Không kể phát đề)")
    format_run(r7, italic=True)

    # Đường kẻ nét đứt xám
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_before = Pt(4)
    p_line.paragraph_format.space_after = Pt(6)
    r_dash = p_line.add_run("---------------------------------------------------------------------------------------------------")
    format_run(r_dash, color=COLOR_GRAY, size_pt=12)

    # Dòng thông tin học sinh
    p_stu = doc.add_paragraph()
    p_stu.paragraph_format.space_before = Pt(2)
    p_stu.paragraph_format.space_after = Pt(8)
    r_stu = p_stu.add_run(
        "Họ và tên học sinh: ......................................................................... "
        "Lớp: ................ Số báo danh: ................"
    )
    format_run(r_stu, italic=True)

    # Tiêu đề PHẦN 1
    p_sec1 = doc.add_paragraph()
    p_sec1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sec1.paragraph_format.space_before = Pt(6)
    p_sec1.paragraph_format.space_after = Pt(8)
    p_sec1.paragraph_format.keep_with_next = True
    r_sec1 = p_sec1.add_run("PHẦN 1: ĐỀ KIỂM TRA")
    format_run(r_sec1, bold=True, color=COLOR_NAVY)


def add_part1_mcq(doc, mcq_questions):
    p_h = doc.add_paragraph()
    p_h.paragraph_format.space_before = Pt(8)
    p_h.paragraph_format.space_after = Pt(2)
    p_h.paragraph_format.keep_with_next = True
    r_h = p_h.add_run("I. TRẮC NGHIỆM NHIỀU LỰA CHỌN (4,0 điểm)")
    format_run(r_h, bold=True)

    p_guide = doc.add_paragraph()
    p_guide.paragraph_format.space_after = Pt(6)
    p_guide.paragraph_format.keep_with_next = True
    r_guide = p_guide.add_run("Thí sinh chọn một phương án trả lời đúng duy nhất trong mỗi câu hỏi từ Câu 1 đến Câu 16.")
    format_run(r_guide, italic=True)

    for idx, q in enumerate(mcq_questions, start=1):
        if q.get('passage'):
            p_pas = doc.add_paragraph()
            p_pas.paragraph_format.space_before = Pt(6)
            p_pas.paragraph_format.space_after = Pt(2)
            p_pas.paragraph_format.keep_with_next = True
            r_pas = p_pas.add_run(f"Đoạn thông tin: {q['passage']}")
            format_run(r_pas, italic=True)

        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.keep_with_next = True
        r_qnum = p_q.add_run(f"Câu {idx}. ")
        format_run(r_qnum, bold=True)
        r_qtxt = p_q.add_run(q['question'])
        format_run(r_qtxt, bold=False)

        opts = q['options']  # ['A. ...', 'B. ...', 'C. ...', 'D. ...']
        wA = measure_text_cm(opts[0])
        wB = measure_text_cm(opts[1])
        wC = measure_text_cm(opts[2])
        wD = measure_text_cm(opts[3])
        maxW = max(wA, wB, wC, wD)
        totalW = wA + wB + wC + wD

        # CÁCH 1: 4 đáp án trên 1 dòng
        if (maxW <= 3.0 and totalW <= 11.0) or (all(len(o) <= 15 for o in opts) and totalW <= 11.5):
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.left_indent = Cm(0.75)
            p_opt.paragraph_format.space_after = Pt(4)
            p_opt.paragraph_format.keep_with_next = True if idx < len(mcq_questions) else False
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(4.69), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(8.63), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(12.57), WD_TAB_ALIGNMENT.LEFT)

            r_lA = p_opt.add_run(opts[0][:3])
            format_run(r_lA, bold=True)
            r_tA = p_opt.add_run(opts[0][3:] + "\t")
            format_run(r_tA)

            r_lB = p_opt.add_run(opts[1][:3])
            format_run(r_lB, bold=True)
            r_tB = p_opt.add_run(opts[1][3:] + "\t")
            format_run(r_tB)

            r_lC = p_opt.add_run(opts[2][:3])
            format_run(r_lC, bold=True)
            r_tC = p_opt.add_run(opts[2][3:] + "\t")
            format_run(r_tC)

            r_lD = p_opt.add_run(opts[3][:3])
            format_run(r_lD, bold=True)
            r_tD = p_opt.add_run(opts[3][3:])
            format_run(r_tD)

        # CÁCH 2: Xếp 2 dòng x 2 đáp án (A-B; C-D), B và D căn tại 7.5 cm
        elif max(wA, wC) <= 6.45 and max(wB, wD) <= 9.0:
            p_row1 = doc.add_paragraph()
            p_row1.paragraph_format.left_indent = Cm(0.75)
            p_row1.paragraph_format.space_after = Pt(1)
            p_row1.paragraph_format.keep_with_next = True
            p_row1.paragraph_format.tab_stops.add_tab_stop(Cm(7.5), WD_TAB_ALIGNMENT.LEFT)

            r_lA = p_row1.add_run(opts[0][:3])
            format_run(r_lA, bold=True)
            r_tA = p_row1.add_run(opts[0][3:] + "\t")
            format_run(r_tA)

            r_lB = p_row1.add_run(opts[1][:3])
            format_run(r_lB, bold=True)
            r_tB = p_row1.add_run(opts[1][3:])
            format_run(r_tB)

            p_row2 = doc.add_paragraph()
            p_row2.paragraph_format.left_indent = Cm(0.75)
            p_row2.paragraph_format.space_after = Pt(4)
            p_row2.paragraph_format.keep_with_next = True if idx < len(mcq_questions) else False
            p_row2.paragraph_format.tab_stops.add_tab_stop(Cm(7.5), WD_TAB_ALIGNMENT.LEFT)

            r_lC = p_row2.add_run(opts[2][:3])
            format_run(r_lC, bold=True)
            r_tC = p_row2.add_run(opts[2][3:] + "\t")
            format_run(r_tC)

            r_lD = p_row2.add_run(opts[3][:3])
            format_run(r_lD, bold=True)
            r_tD = p_row2.add_run(opts[3][3:])
            format_run(r_tD)

        # CÁCH 3: Mỗi đáp án một dòng
        else:
            for o_idx, opt in enumerate(opts):
                p_single = doc.add_paragraph()
                p_single.paragraph_format.left_indent = Cm(0.75)
                p_single.paragraph_format.space_after = Pt(1 if o_idx < 3 else 4)
                p_single.paragraph_format.keep_with_next = True if o_idx < 3 else (idx < len(mcq_questions))

                r_lbl = p_single.add_run(opt[:3])
                format_run(r_lbl, bold=True)
                r_txt = p_single.add_run(opt[3:])
                format_run(r_txt)


def add_part2_true_false(doc, tf_questions):
    p_h = doc.add_paragraph()
    p_h.paragraph_format.space_before = Pt(8)
    p_h.paragraph_format.space_after = Pt(2)
    p_h.paragraph_format.keep_with_next = True
    r_h = p_h.add_run("II. TRẮC NGHIỆM ĐÚNG / SAI (3,0 điểm)")
    format_run(r_h, bold=True)

    p_guide = doc.add_paragraph()
    p_guide.paragraph_format.space_after = Pt(6)
    p_guide.paragraph_format.keep_with_next = True
    r_guide = p_guide.add_run("Thí sinh trả lời từ Câu 1 đến Câu 3. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")
    format_run(r_guide, italic=True)

    for idx, q in enumerate(tf_questions, start=1):
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.keep_with_next = True
        r_qnum = p_q.add_run(f"Câu {idx}. ")
        format_run(r_qnum, bold=True)
        r_qtxt = p_q.add_run(q['context'])
        format_run(r_qtxt)

        for s_idx, item in enumerate(q['items']):
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.left_indent = Cm(0.75)
            p_sub.paragraph_format.space_after = Pt(2 if s_idx < 3 else 4)
            p_sub.paragraph_format.keep_with_next = True if s_idx < 3 else (idx < len(tf_questions))
            r_lbl = p_sub.add_run(f"{item['label']}) ")
            format_run(r_lbl, bold=True)
            r_stmt = p_sub.add_run(item['statement'])
            format_run(r_stmt)


def add_part3_essay(doc, essay_questions):
    p_h = doc.add_paragraph()
    p_h.paragraph_format.space_before = Pt(8)
    p_h.paragraph_format.space_after = Pt(2)
    p_h.paragraph_format.keep_with_next = True
    r_h = p_h.add_run("III. TỰ LUẬN (3,0 điểm)")
    format_run(r_h, bold=True)

    p_guide = doc.add_paragraph()
    p_guide.paragraph_format.space_after = Pt(6)
    p_guide.paragraph_format.keep_with_next = True
    r_guide = p_guide.add_run("Thí sinh trình bày chi tiết lời giải các câu hỏi sau vào giấy làm bài.")
    format_run(r_guide, italic=True)

    for idx, q in enumerate(essay_questions, start=1):
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.keep_with_next = True
        r_qnum = p_q.add_run(f"Câu {idx} (1,0 điểm). ")
        format_run(r_qnum, bold=True)
        r_qtxt = p_q.add_run(q['context'])
        format_run(r_qtxt)

        p_a = doc.add_paragraph()
        p_a.paragraph_format.space_before = Pt(1)
        p_a.paragraph_format.space_after = Pt(1)
        p_a.paragraph_format.keep_with_next = True
        r_albl = p_a.add_run("a) (0,5 điểm) ")
        format_run(r_albl, bold=True)
        r_atxt = p_a.add_run(q['part_a'])
        format_run(r_atxt)

        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = Pt(1)
        p_b.paragraph_format.space_after = Pt(4)
        p_b.paragraph_format.keep_with_next = True if idx < len(essay_questions) else False
        r_blbl = p_b.add_run("b) (0,5 điểm) ")
        format_run(r_blbl, bold=True)
        r_btxt = p_b.add_run(q['part_b'])
        format_run(r_btxt)

    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_end.paragraph_format.space_before = Pt(14)
    p_end.paragraph_format.space_after = Pt(2)
    r_end = p_end.add_run("------------------- HẾT -------------------")
    format_run(r_end, bold=True, italic=True)

    p_note = doc.add_paragraph()
    p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_note.paragraph_format.space_after = Pt(6)
    r_note = p_note.add_run("(Cán bộ coi thi không giải thích gì thêm)")
    format_run(r_note, italic=True)


def add_part2_answers_and_rubric(doc, mcq_answers, tf_answers, essay_rubrics):
    doc.add_page_break()

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(8)
    p_title.paragraph_format.keep_with_next = True
    r_t = p_title.add_run("PHẦN 2: ĐÁP ÁN - HƯỚNG DẪN CHẤM")
    format_run(r_t, bold=True, color=COLOR_NAVY)

    # I. BẢNG ĐÁP ÁN TRẮC NGHIỆM
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(6)
    p_h1.paragraph_format.space_after = Pt(2)
    p_h1.paragraph_format.keep_with_next = True
    r_h1 = p_h1.add_run("I. ĐÁP ÁN PHẦN TRẮC NGHIỆM NHIỀU LỰA CHỌN (4,0 điểm)")
    format_run(r_h1, bold=True)

    p_g1 = doc.add_paragraph()
    p_g1.paragraph_format.space_after = Pt(6)
    p_g1.paragraph_format.keep_with_next = True
    r_g1 = p_g1.add_run("Mỗi câu trả lời đúng được 0,25 điểm. Tổng cộng 16 câu = 4,0 điểm.")
    format_run(r_g1, italic=True)

    tbl_mcq = doc.add_table(rows=4, cols=9)
    tbl_mcq.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_mcq.autofit = False
    set_table_borders(tbl_mcq, border_color=HEX_BORDER, border_size="4")

    col_w_label = Cm(2.1)
    col_w_data = Cm(1.8)

    # Hàng 0: Câu 1-8
    row0 = tbl_mcq.rows[0]
    set_row_cant_split(row0)
    c = row0.cells[0]
    c.width = col_w_label
    set_cell_shading(c, HEX_BG_GRAY)
    set_cell_margins(c, 3, 3, 2, 2)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_run(p.add_run("Câu"), bold=True)

    for i in range(1, 9):
        c = row0.cells[i]
        c.width = col_w_data
        set_cell_shading(c, HEX_BG_GRAY)
        set_cell_margins(c, 3, 3, 2, 2)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_run(p.add_run(str(i)), bold=True)

    # Hàng 1: Đáp án 1-8
    row1 = tbl_mcq.rows[1]
    set_row_cant_split(row1)
    c = row1.cells[0]
    c.width = col_w_label
    set_cell_shading(c, HEX_BG_HEADER_LIGHT)
    set_cell_margins(c, 3, 3, 2, 2)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_run(p.add_run("Đáp án"), bold=True, color=COLOR_RED_LABEL)

    for i in range(1, 9):
        c = row1.cells[i]
        c.width = col_w_data
        set_cell_margins(c, 3, 3, 2, 2)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_run(p.add_run(mcq_answers[i - 1]), bold=True, color=COLOR_RED_LABEL)

    # Hàng 2: Câu 9-16
    row2 = tbl_mcq.rows[2]
    set_row_cant_split(row2)
    c = row2.cells[0]
    c.width = col_w_label
    set_cell_shading(c, HEX_BG_GRAY)
    set_cell_margins(c, 3, 3, 2, 2)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_run(p.add_run("Câu"), bold=True)

    for i in range(1, 9):
        q_idx = 8 + i
        c = row2.cells[i]
        c.width = col_w_data
        set_cell_shading(c, HEX_BG_GRAY)
        set_cell_margins(c, 3, 3, 2, 2)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_run(p.add_run(str(q_idx)), bold=True)

    # Hàng 3: Đáp án 9-16
    row3 = tbl_mcq.rows[3]
    set_row_cant_split(row3)
    c = row3.cells[0]
    c.width = col_w_label
    set_cell_shading(c, HEX_BG_HEADER_LIGHT)
    set_cell_margins(c, 3, 3, 2, 2)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_run(p.add_run("Đáp án"), bold=True, color=COLOR_RED_LABEL)

    for i in range(1, 9):
        q_idx = 8 + i
        c = row3.cells[i]
        c.width = col_w_data
        set_cell_margins(c, 3, 3, 2, 2)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_run(p.add_run(mcq_answers[q_idx - 1]), bold=True, color=COLOR_RED_LABEL)

    # II. BẢNG ĐÁP ÁN ĐÚNG / SAI
    p_h2 = doc.add_paragraph()
    p_h2.paragraph_format.space_before = Pt(10)
    p_h2.paragraph_format.space_after = Pt(2)
    p_h2.paragraph_format.keep_with_next = True
    r_h2 = p_h2.add_run("II. ĐÁP ÁN PHẦN TRẮC NGHIỆM ĐÚNG / SAI (3,0 điểm)")
    format_run(r_h2, bold=True)

    rules = [
        "Quy tắc tính điểm theo quy định Bộ GD&ĐT cho mỗi câu lớn (1,0 điểm/câu):",
        "- Trả lời đúng 01 ý: được 0,10 điểm.",
        "- Trả lời đúng 02 ý: được 0,25 điểm.",
        "- Trả lời đúng 03 ý: được 0,50 điểm.",
        "- Trả lời đúng cả 04 ý: được 1,00 điểm.",
    ]
    for r_idx, rule in enumerate(rules):
        p_r = doc.add_paragraph()
        p_r.paragraph_format.space_after = Pt(1 if r_idx < 4 else 6)
        p_r.paragraph_format.keep_with_next = True
        r_txt = p_r.add_run(rule)
        format_run(r_txt, italic=True)

    tbl_tf = doc.add_table(rows=13, cols=4)
    tbl_tf.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_tf.autofit = False
    set_table_borders(tbl_tf, border_color=HEX_BORDER, border_size="4")

    row_hdr = tbl_tf.rows[0]
    set_row_cant_split(row_hdr)
    set_row_header(row_hdr)
    hdr_titles = ["Câu hỏi", "Lệnh hỏi", "Đáp án", "Hướng dẫn giải thích ngắn gọn"]
    hdr_widths = [Cm(2.8), Cm(1.8), Cm(2.4), Cm(9.5)]
    for idx, (title, w) in enumerate(zip(hdr_titles, hdr_widths)):
        c = row_hdr.cells[idx]
        c.width = w
        set_cell_shading(c, HEX_BG_BLUE_LIGHT)
        set_cell_margins(c, 4, 4, 3, 3)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_run(p.add_run(title), bold=True)

    current_row_idx = 1
    for q_idx, q in enumerate(tf_answers, start=1):
        start_row = current_row_idx
        for s_idx, item in enumerate(q['items']):
            row = tbl_tf.rows[current_row_idx]
            set_row_cant_split(row)

            c0 = row.cells[0]
            c0.width = Cm(2.8)

            c1 = row.cells[1]
            c1.width = Cm(1.8)
            set_cell_margins(c1, 3, 3, 2, 2)
            c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p1 = c1.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            format_run(p1.add_run(f"{item['label']})"), bold=True)

            c2 = row.cells[2]
            c2.width = Cm(2.4)
            set_cell_margins(c2, 3, 3, 2, 2)
            c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p2 = c2.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            is_true = item['answer'] == 'ĐÚNG'
            ans_color = COLOR_GREEN if is_true else COLOR_RED_SAI
            format_run(p2.add_run(item['answer']), bold=True, color=ans_color)

            c3 = row.cells[3]
            c3.width = Cm(9.5)
            set_cell_margins(c3, 3, 3, 4, 4)
            c3.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p3 = c3.paragraphs[0]
            p3.alignment = WD_ALIGN_PARAGRAPH.LEFT
            format_run(p3.add_run(item['explanation']))

            current_row_idx += 1

        c_top = tbl_tf.cell(start_row, 0)
        c_bottom = tbl_tf.cell(start_row + 3, 0)
        c_merged = c_top.merge(c_bottom)
        c_merged.width = Cm(2.8)
        set_cell_margins(c_merged, 4, 4, 2, 2)
        c_merged.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p0 = c_merged.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_run(p0.add_run(f"Câu {q_idx}\n"), bold=True)
        format_run(p0.add_run("(1,0 điểm)"), bold=True)

    # III. ĐÁP ÁN VÀ THANG ĐIỂM PHẦN TỰ LUẬN
    p_h3 = doc.add_paragraph()
    p_h3.paragraph_format.space_before = Pt(10)
    p_h3.paragraph_format.space_after = Pt(6)
    p_h3.paragraph_format.keep_with_next = True
    r_h3 = p_h3.add_run("III. ĐÁP ÁN VÀ THANG ĐIỂM PHẦN TỰ LUẬN (3,0 điểm)")
    format_run(r_h3, bold=True)

    for q_idx, q in enumerate(essay_rubrics, start=1):
        tbl_e = doc.add_table(rows=5, cols=2)
        tbl_e.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_e.autofit = False
        set_table_borders(tbl_e, border_color=HEX_BORDER, border_size="4")

        row_title = tbl_e.rows[0]
        set_row_cant_split(row_title)
        c_t0 = row_title.cells[0]
        c_t1 = row_title.cells[1]
        c_tm = c_t0.merge(c_t1)
        c_tm.width = Cm(16.5)
        set_cell_shading(c_tm, HEX_BG_HEADER_LIGHT)
        set_cell_margins(c_tm, 4, 4, 5, 5)
        p_tm = c_tm.paragraphs[0]
        format_run(p_tm.add_run(f"Câu {q_idx} (1,0 điểm):"), bold=True)

        for s_idx, step in enumerate(q['steps']):
            row_s = tbl_e.rows[s_idx + 1]
            set_row_cant_split(row_s)

            c_content = row_s.cells[0]
            c_content.width = Cm(13.5)
            set_cell_margins(c_content, 3, 3, 5, 5)
            c_content.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p_c = c_content.paragraphs[0]
            format_run(p_c.add_run(f"• {step['content']}"))

            c_pt = row_s.cells[1]
            c_pt.width = Cm(3.0)
            set_cell_margins(c_pt, 3, 3, 2, 2)
            c_pt.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p_pt = c_pt.paragraphs[0]
            p_pt.alignment = WD_ALIGN_PARAGRAPH.CENTER
            format_run(p_pt.add_run("0,25 điểm"), bold=True)

        if q_idx < len(essay_rubrics):
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_after = Pt(4)


# ==============================================================================
# DỮ LIỆU CHUYÊN MÔN 3 KHỐI LỚP (GDPT 2018 - KẾT NỐI TRI THỨC VỚI CUỘC SỐNG)
# ==============================================================================

EXAM_DATA_GRADE_7 = {
    'school_info': {
        'dept': 'PHÒNG GD&ĐT HUYỆN LONG PHÚ',
        'school': 'TRƯỜNG THCS LONG PHÚ',
    },
    'exam_info': {
        'title': 'ĐỀ KIỂM TRA ĐỊNH KỲ GIỮA HỌC KỲ I',
        'year': '2026 - 2027',
        'subject': 'KHOA HỌC TỰ NHIÊN',
        'grade': '7',
        'code': '701',
        'duration': 45,
    },
    'mcq': [
        {
            'question': 'Phương pháp tìm hiểu tự nhiên được thực hiện qua bao nhiêu bước cơ bản?',
            'options': ['A. 3 bước.', 'B. 4 bước.', 'C. 5 bước.', 'D. 6 bước.'],
            'answer': 'C'
        },
        {
            'question': 'Dụng cụ nào sau đây thường được dùng để quan sát những vật thể có kích thước rất nhỏ mà mắt thường không thể nhìn thấy được?',
            'options': ['A. Kính hiển vi quang học.', 'B. Kính lúp cầm tay.', 'C. Thước kẹp cơ khí.', 'D. Ống nhòm du lịch.'],
            'answer': 'A'
        },
        {
            'question': 'Hoạt động nào sau đây KHÔNG thuộc các kĩ năng tiến trình tìm hiểu tự nhiên của học sinh?',
            'options': ['A. Kĩ năng quan sát và đặt câu hỏi.', 'B. Kĩ năng phân loại và đo lường.', 'C. Kĩ năng liên kết và dự đoán.', 'D. Kĩ năng phán đoán cảm tính.'],
            'answer': 'D'
        },
        {
            'question': 'Khi sử dụng kính lúp cầm tay để quan sát gân lá cây, ta cần điều chỉnh khoảng cách giữa kính lúp và mẫu vật nhằm mục đích',
            'options': ['A. làm đổi màu sắc của mẫu vật.', 'B. nhìn rõ ảnh phóng to của vật.', 'C. làm cho mẫu vật phát sáng.', 'D. biến hình ảnh thu nhỏ lại.'],
            'answer': 'B'
        },
        {
            'question': 'Để kiểm chứng giả thuyết "Nhiệt độ nước càng cao thì đường tan càng nhanh", học sinh cần tiến hành bước nào sau đây?',
            'options': ['A. Viết ngay kết luận vào vở ghi.', 'B. Làm thí nghiệm đối chứng nhiệt độ.', 'C. Đọc lại mục lục sách giáo khoa.', 'D. Chọn màu sắc cho bài trình chiếu.'],
            'answer': 'B'
        },
        {
            'question': 'Biển cảnh báo có hình tam giác viền đen, nền vàng và ngọn lửa màu đen mang ý nghĩa cảnh báo điều gì trong phòng thực hành?',
            'options': ['A. Hóa chất ăn mòn kim loại.', 'B. Nguồn điện cao thế nguy hiểm.', 'C. Hóa chất dễ cháy.', 'D. Chất độc hại sinh học.'],
            'answer': 'C'
        },
        {
            'question': 'Hạt nhân của nguyên tử được cấu tạo từ những loại hạt nào sau đây?',
            'options': ['A. Proton và electron.', 'B. Proton và neutron.', 'C. Neutron và electron.', 'D. Chỉ có hạt electron.'],
            'answer': 'B'
        },
        {
            'question': 'Trong nguyên tử, loại hạt mang điện tích âm và chuyển động xung quanh hạt nhân là',
            'options': ['A. electron.', 'B. proton.', 'C. neutron.', 'D. hạt nhân.'],
            'answer': 'A'
        },
        {
            'question': 'Nguyên tử của nguyên tố nhôm (Aluminium) có 13 proton trong hạt nhân. Số electron chuyển động ở lớp vỏ nguyên tử này là',
            'options': ['A. 11.', 'B. 12.', 'C. 13.', 'D. 14.'],
            'answer': 'C'
        },
        {
            'question': 'Trong bảng tuần hoàn, các nguyên tố hóa học được sắp xếp theo quy tắc chiều tăng dần của',
            'options': ['A. khối lượng nguyên tử.', 'B. bán kính nguyên tử.', 'C. số electron lớp ngoài cùng.', 'D. điện tích hạt nhân nguyên tử.'],
            'answer': 'D'
        },
        {
            'question': 'Tập hợp các nguyên tố hóa học có cùng số lớp electron trong nguyên tử được xếp vào cùng một',
            'options': ['A. chu kì.', 'B. nhóm A.', 'C. nhóm B.', 'D. ô nguyên tố.'],
            'answer': 'A'
        },
        {
            'question': 'Đại lượng vật lí nào sau đây cho biết mức độ chuyển động nhanh hay chậm của một vật thể?',
            'options': ['A. Quãng đường đi được.', 'B. Thời gian chuyển động.', 'C. Tốc độ chuyển động.', 'D. Trọng lượng của vật.'],
            'answer': 'C'
        },
        {
            'question': 'Biểu thức xác định tốc độ chuyển động v của một vật đi được quãng đường s trong thời gian t là',
            'options': ['A. v = s / t.', 'B. v = s . t.', 'C. v = t / s.', 'D. v = 2s / t.'],
            'answer': 'A'
        },
        {
            'question': 'Một ô tô đang lưu thông trên đường với tốc độ 72\u00A0km/h. Tốc độ của ô tô tính theo đơn vị m/s là',
            'options': ['A. 15\u00A0m/s.', 'B. 20\u00A0m/s.', 'C. 25\u00A0m/s.', 'D. 30\u00A0m/s.'],
            'answer': 'B'
        },
        {
            'question': 'Một bạn học sinh đạp xe đều trên quãng đường thẳng dài 600\u00A0m mất thời gian đúng 2\u00A0phút. Tốc độ chuyển động của bạn học sinh là',
            'options': ['A. 2\u00A0m/s.', 'B. 3\u00A0m/s.', 'C. 4\u00A0m/s.', 'D. 5\u00A0m/s.'],
            'answer': 'D'
        },
        {
            'question': 'Thiết bị nào sau đây thường được cảnh sát giao thông sử dụng để kiểm tra tốc độ phương tiện trên đường?',
            'options': ['A. Đồng hồ bấm giây hiện số.', 'B. Thước dây cuộn đo chiều dài.', 'C. Cổng quang điện cảm ứng.', 'D. Súng bắn tốc độ nghiệp vụ.'],
            'answer': 'D'
        },
    ],
    'tf': [
        {
            'context': 'Trong giờ thực hành môn Khoa học tự nhiên 7, nhóm học sinh thực hiện phép đo khối lượng mẫu sỏi nhỏ và làm tiêu bản quan sát tế bào biểu bì vảy hành dưới kính hiển vi quang học.',
            'items': [
                {'label': 'a', 'statement': 'Học sinh có thể sử dụng cân điện tử để xác định chính xác khối lượng của mẫu sỏi nhỏ.', 'answer': 'ĐÚNG', 'explanation': 'Cân điện tử là dụng cụ thích hợp để đo khối lượng vật nhỏ với độ chính xác cao.'},
                {'label': 'b', 'statement': 'Khi quan sát tế bào biểu bì hành, ta đặt lam kính chứa mẫu tự do trên bàn kính mà không cần kẹp giữ.', 'answer': 'SAI', 'explanation': 'Bắt buộc phải dùng kẹp giữ tiêu bản để cố định vị trí mẫu vật quan sát trên bàn kính.'},
                {'label': 'c', 'statement': 'Học sinh tiến hành đo khối lượng mẫu sỏi 3 lần rồi lấy giá trị trung bình để giảm thiểu sai số phép đo.', 'answer': 'ĐÚNG', 'explanation': 'Thực hiện lặp lại phép đo nhiều lần giúp hạn chế sai số ngẫu nhiên trong thực nghiệm.'},
                {'label': 'd', 'statement': 'Khi di chuyển kính hiển vi quang học, chỉ cần dùng một tay nắm lấy ống kính và kéo rê trên mặt bàn.', 'answer': 'SAI', 'explanation': 'Quy tắc an toàn: Dùng hai tay, một tay đỡ chân đế và một tay giữ chắc thân kính.'},
            ]
        },
        {
            'context': 'Cho biết nguyên tử của nguyên tố X có tổng số hạt cơ bản (proton, neutron, electron) là 40, trong đó số hạt mang điện nhiều hơn số hạt không mang điện là 12 hạt.',
            'items': [
                {'label': 'a', 'statement': 'Trong nguyên tử X, số hạt proton bằng số hạt electron và có giá trị bằng 13.', 'answer': 'ĐÚNG', 'explanation': 'Hệ phương trình: 2p + n = 40 và 2p - n = 12 => 4p = 52 => p = e = 13 (Aluminium).'},
                {'label': 'b', 'statement': 'Số lượng hạt neutron nằm trong hạt nhân của nguyên tử X là 14 hạt.', 'answer': 'ĐÚNG', 'explanation': 'Từ hệ phương trình ta tính được: n = 40 - 2p = 40 - 26 = 14 hạt neutron.'},
                {'label': 'c', 'statement': 'Khối lượng nguyên tử X xấp xỉ bằng khối lượng của hạt nhân và bằng 27\u00A0amu.', 'answer': 'ĐÚNG', 'explanation': 'Khối lượng nguyên tử = số p + số n = 13 + 14 = 27 amu (bỏ qua khối lượng electron rất nhỏ).'},
                {'label': 'd', 'statement': 'Nguyên tố X thuộc chu kì 2 và nhóm IIA trong bảng tuần hoàn các nguyên tố hóa học.', 'answer': 'SAI', 'explanation': 'Aluminium (p = 13) có 3 lớp electron và 3e lớp ngoài cùng nên thuộc chu kì 3, nhóm IIIA.'},
            ]
        },
        {
            'context': 'Một vận động viên điền kinh thực hiện phần thi chạy cự li 100\u00A0m trên đường thẳng. Trong 4\u00A0giây đầu tiên người đó chạy được 32\u00A0m; quãng đường 68\u00A0m còn lại chạy hết trong thời gian 6\u00A0giây tiếp theo.',
            'items': [
                {'label': 'a', 'statement': 'Tốc độ trung bình của vận động viên trong 4\u00A0giây đầu tiên đạt giá trị là 8\u00A0m/s.', 'answer': 'ĐÚNG', 'explanation': 'Tốc độ 4 giây đầu: v1 = s1 / t1 = 32 / 4 = 8 m/s.'},
                {'label': 'b', 'statement': 'Tốc độ trung bình của vận động viên tính trên toàn bộ cự li 100\u00A0m là 10\u00A0m/s.', 'answer': 'ĐÚNG', 'explanation': 'Tốc độ cả quãng đường: v_tb = s_tong / t_tong = 100 / (4 + 6) = 100 / 10 = 10 m/s.'},
                {'label': 'c', 'statement': 'Chuyển động của vận động viên trên cả quãng đường 100\u00A0m là chuyển động đều vì tốc độ không đổi.', 'answer': 'SAI', 'explanation': 'Tốc độ ở hai giai đoạn khác nhau (8 m/s và 68/6 ≈ 11,33 m/s) nên là chuyển động không đều.'},
                {'label': 'd', 'statement': 'Đơn vị đo tốc độ chính thức, hợp pháp trong Hệ đo lường quốc tế SI là kilômét trên giờ (km/h).', 'answer': 'SAI', 'explanation': 'Đơn vị đo tốc độ chuẩn trong hệ đo lường quốc tế SI là mét trên giây (m/s).'},
            ]
        },
    ],
    'essay': [
        {
            'context': 'Trong môn Khoa học tự nhiên, phương pháp tìm hiểu tự nhiên là công cụ quan trọng giúp học sinh khám phá thế giới xung quanh.',
            'part_a': 'Trình bày đầy đủ 5 bước cơ bản của tiến trình phương pháp tìm hiểu tự nhiên.',
            'part_b': 'Đề xuất 01 giả thuyết khoa học và nêu cách bố trí thí nghiệm đơn giản để kiểm chứng: "Ánh sáng ảnh hưởng đến sự sinh trưởng của cây non".',
            'steps': [
                {'content': 'Nêu đúng 5 bước: (1) Đề xuất vấn đề; (2) Đưa ra giả thuyết khoa học; (3) Lập kế hoạch kiểm tra; (4) Thực hiện kế hoạch; (5) Viết báo cáo và thảo luận.'},
                {'content': 'Giải thích được vai trò: Việc thực nghiệm kiểm chứng giúp khẳng định giả thuyết đặt ra là đúng hay sai.'},
                {'content': 'Đề xuất giả thuyết phù hợp: Cây non trồng ở nơi có đủ ánh sáng mặt trời sẽ sinh trưởng phát triển nhanh và xanh tốt hơn cây trong bóng tối.'},
                {'content': 'Bố trí thí nghiệm đối chứng: Trồng 2 chậu cây giống nhau (cùng đất, tưới cùng lượng nước), 1 chậu để ngoài sáng, 1 chậu đặt trong hộp tối kín; quan sát sau 5 - 7 ngày.'},
            ]
        },
        {
            'context': 'Nguyên tử của nguyên tố Calcium (kí hiệu hóa học: Ca) có điện tích hạt nhân là +20.',
            'part_a': 'Xác định số proton, số electron trong nguyên tử Calcium và mô tả sự phân bố các electron vào các lớp vỏ (từ trong ra ngoài).',
            'part_b': 'Dựa vào cấu tạo nguyên tử, hãy xác định vị trí của Calcium trong bảng tuần hoàn (ô nguyên tố, chu kì, nhóm) và cho biết đây là kim loại hay phi kim.',
            'steps': [
                {'content': 'Số proton = 20; vì nguyên tử trung hòa về điện nên số electron = 20.'},
                {'content': 'Mô tả phân bố electron vào 4 lớp vỏ: Lớp 1 có 2e, Lớp 2 có 8e, Lớp 3 có 8e, Lớp 4 (lớp ngoài cùng) có 2e.'},
                {'content': 'Xác định vị trí: Ô số 20 (vì có 20 proton), Chu kì 4 (vì có 4 lớp electron), Nhóm IIA (vì có 2 electron ở lớp ngoài cùng).'},
                {'content': 'Kết luận: Nguyên tố Calcium có 2 electron ở lớp vỏ ngoài cùng nên là nguyên tố KIM LOẠI điển hình.'},
            ]
        },
        {
            'context': 'Hai thành phố A và B cách nhau 120\u00A0km. Lúc 7\u00A0giờ sáng, ô tô khởi hành từ A về B với tốc độ không đổi 60\u00A0km/h. Cùng lúc đó, xe máy khởi hành từ B về A với tốc độ không đổi 40\u00A0km/h trên cùng tuyến đường thẳng.',
            'part_a': 'Viết công thức tính quãng đường theo thời gian và tính khoảng thời gian từ lúc xuất phát cho đến khi hai xe gặp nhau.',
            'part_b': 'Xác định thời điểm hai xe gặp nhau và vị trí điểm gặp nhau cách thành phố A bao nhiêu kilômét?',
            'steps': [
                {'content': 'Viết công thức s = v.t; tổng quãng đường hai xe đi được khi gặp nhau: s1 + s2 = s_AB = 120 km.'},
                {'content': 'Tính thời gian chuyển động cho đến khi gặp nhau: t = s_AB / (v1 + v2) = 120 / (60 + 40) = 1,2 giờ (tức 1 giờ 12 phút).'},
                {'content': 'Thời điểm hai xe gặp nhau: 7 giờ 00 phút + 1 giờ 12 phút = 8 giờ 12 phút sáng.'},
                {'content': 'Vị trí điểm gặp nhau cách điểm xuất phát A: s1 = v1 . t = 60 . 1,2 = 72 km.'},
            ]
        },
    ]
}


EXAM_DATA_GRADE_8 = {
    'school_info': {
        'dept': 'PHÒNG GD&ĐT HUYỆN LONG PHÚ',
        'school': 'TRƯỜNG THCS LONG PHÚ',
    },
    'exam_info': {
        'title': 'ĐỀ KIỂM TRA ĐỊNH KỲ GIỮA HỌC KỲ I',
        'year': '2026 - 2027',
        'subject': 'KHOA HỌC TỰ NHIÊN',
        'grade': '8',
        'code': '801',
        'duration': 45,
    },
    'mcq': [
        {
            'question': 'Dụng cụ nào sau đây thường dùng để lấy một lượng chất lỏng có thể tích chính xác trong thí nghiệm hoá học?',
            'options': ['A. Pipette có chia vạch.', 'B. Cốc thủy tinh chịu nhiệt.', 'C. Bình tam giác đáy phẳng.', 'D. Bát sứ nung mẫu.'],
            'answer': 'A'
        },
        {
            'question': 'Khi cần pha loãng dung dịch acid đặc (H2SO4 đặc) trong phòng thực hành, thao tác nào sau đây là an toàn và ĐÚNG kĩ thuật?',
            'options': ['A. Đổ nhanh nước lạnh vào acid đặc.', 'B. Rót từ từ acid đặc vào nước.', 'C. Dùng tay khuấy đều hỗn hợp.', 'D. Đổ acid và nước cùng một lúc.'],
            'answer': 'B'
        },
        {
            'question': 'Để dập tắt đám cháy nhỏ bắt nguồn từ xăng dầu trong phòng thí nghiệm, phương pháp nào sau đây là KHÔNG được sử dụng?',
            'options': ['A. Dùng bình chữa cháy khí CO2.', 'B. Phủ cát khô trùm lên ngọn lửa.', 'C. Dội nhiều nước vào đám cháy.', 'D. Dùng chăn chiên dày trùm kín.'],
            'answer': 'C'
        },
        {
            'question': 'Thiết bị nào sau đây dùng để xác định chính xác khối lượng của mẫu hoá chất rắn trong phòng thực hành KHTN 8?',
            'options': ['A. Nhiệt kế thủy ngân.', 'B. Cân điện tử hiện số.', 'C. Đồng hồ bấm giây.', 'D. Ống đong chia độ.'],
            'answer': 'B'
        },
        {
            'question': 'Kí hiệu cảnh báo hình đầu lâu xương chéo in trên nhãn bao bì chai lọ hoá chất mang ý nghĩa cảnh báo điều gì?',
            'options': ['A. Hóa chất dễ cháy nổ.', 'B. Hóa chất ăn mòn kim loại.', 'C. Hóa chất độc hại nguy hiểm.', 'D. Chất phát ra tia phóng xạ.'],
            'answer': 'C'
        },
        {
            'question': 'Khi tiến hành đun nóng chất lỏng trong ống nghiệm bằng ngọn lửa đèn cồn, miệng ống nghiệm phải được hướng về phía',
            'options': ['A. người làm thí nghiệm.', 'B. bạn học ngồi đối diện.', 'C. nơi không có người.', 'D. tủ đựng hóa chất dễ cháy.'],
            'answer': 'C'
        },
        {
            'question': 'Hiện tượng nào sau đây là biểu hiện rõ nét của một hiện tượng biến đổi hoá học?',
            'options': ['A. Nước lỏng đông đặc thành đá.', 'B. Đinh sắt để ngoài ẩm bị gỉ.', 'C. Cồn hở miệng bay hơi dần.', 'D. Hòa tan muối ăn vào nước.'],
            'answer': 'B'
        },
        {
            'question': 'Theo định luật bảo toàn khối lượng, trong một phản ứng hoá học, đại lượng nào sau đây luôn được bảo toàn?',
            'options': ['A. Tổng khối lượng các chất.', 'B. Số lượng phân tử các chất.', 'C. Thể tích hỗn hợp phản ứng.', 'D. Cấu trúc liên kết phân tử.'],
            'answer': 'A'
        },
        {
            'question': 'Ở điều kiện chuẩn (nhiệt độ 25 °C và áp suất 1\u00A0bar), thể tích của 1\u00A0mol chất khí bất kì bằng',
            'options': ['A. 22,4\u00A0L.', 'B. 22,7\u00A0L.', 'C. 25,0\u00A0L.', 'D. 24,79\u00A0L.'],
            'answer': 'D'
        },
        {
            'question': 'Khối lượng mol phân tử (M) của hợp chất khí carbon dioxide (công thức phân tử CO2) có giá trị là',
            'options': ['A. 28\u00A0g/mol.', 'B. 32\u00A0g/mol.', 'C. 44\u00A0g/mol.', 'D. 56\u00A0g/mol.'],
            'answer': 'C'
        },
        {
            'question': 'Tỉ khối của khí methane (CH4, M = 16\u00A0g/mol) so với không khí (M_kk xấp xỉ 29\u00A0g/mol) cho biết khí methane',
            'options': ['A. nặng hơn không khí 1,8 lần.', 'B. nhẹ hơn, bằng khoảng 16/29.', 'C. nặng gấp 2 lần không khí.', 'D. nặng bằng khối lượng không khí.'],
            'answer': 'B'
        },
        {
            'question': 'Áp suất được định nghĩa trong vật lí là',
            'options': ['A. độ lớn áp lực trên 1 đv diện tích.', 'B. lực ép song song mặt tiếp xúc.', 'C. tích áp lực và diện tích ép.', 'D. trọng lượng của vật tác dụng.'],
            'answer': 'A'
        },
        {
            'question': 'Đơn vị đo áp suất chính thức trong Hệ đo lường quốc tế SI là',
            'options': ['A. Niutơn (N).', 'B. Jun (J).', 'C. Kilôgam (kg).', 'D. Paxcan (Pa).'],
            'answer': 'D'
        },
        {
            'question': 'Áp suất do chất lỏng tác dụng lên đáy bình chứa phụ thuộc trực tiếp vào những yếu tố nào sau đây?',
            'options': ['A. Diện tích của bề mặt đáy bình.', 'B. Thể tích toàn phần của bình.', 'C. Hình dáng thành của bình chứa.', 'D. Khối lượng riêng và độ sâu.'],
            'answer': 'D'
        },
        {
            'question': 'Một khối kim loại có thể tích 0,002\u00A0m³ và khối lượng đo được là 5,4\u00A0kg. Khối lượng riêng D của kim loại đó là',
            'options': ['A. 2\u00A0700\u00A0kg/m³.', 'B. 1\u00A0350\u00A0kg/m³.', 'C. 5\u00A0400\u00A0kg/m³.', 'D. 10\u00A0800\u00A0kg/m³.'],
            'answer': 'A'
        },
        {
            'question': 'Một vật thể khi được nhúng chìm hoàn toàn trong lòng chất lỏng sẽ chịu tác dụng của lực đẩy Archimedes có hướng',
            'options': ['A. nằm ngang từ trái qua phải.', 'B. xiên góc 45 độ so với đáy.', 'C. thẳng đứng từ trên xuống.', 'D. thẳng đứng từ dưới lên.'],
            'answer': 'D'
        },
    ],
    'tf': [
        {
            'context': 'Để xác định nồng độ dung dịch acid HCl bằng phương pháp chuẩn độ với dung dịch base NaOH chuẩn có chỉ thị phenolphthalein, nhóm học sinh lớp 8 chuẩn bị buret, pipet, bình tam giác và tiến hành thí nghiệm.',
            'items': [
                {'label': 'a', 'statement': 'Học sinh cần tráng buret bằng chính dung dịch chuẩn cần nạp trước khi tiến hành thí nghiệm.', 'answer': 'ĐÚNG', 'explanation': 'Tráng buret bằng dung dịch cần chứa giúp nồng độ dung dịch trong buret không bị thay đổi bởi nước sót lại.'},
                {'label': 'b', 'statement': 'Dung dịch kiềm NaOH có tính nhờn và không ăn mòn nên có thể để dính tự do vào da hoặc mắt.', 'answer': 'SAI', 'explanation': 'NaOH là base mạnh ăn mòn da và phá hủy giác mạc rất nguy hiểm, phải dùng đồ bảo hộ.'},
                {'label': 'c', 'statement': 'Tại điểm tương đương, dung dịch trong bình tam giác chuyển từ không màu sang màu hồng nhạt bền trong 30 giây.', 'answer': 'ĐÚNG', 'explanation': 'Phenolphthalein chuyển màu hồng trong môi trường kiềm nhẹ, đánh dấu kết thúc chuẩn độ.'},
                {'label': 'd', 'statement': 'Sau khi thực hành xong, toàn bộ hoá chất dư được đổ trực tiếp vào cống thoát nước mà không cần thu gom.', 'answer': 'SAI', 'explanation': 'Quy tắc an toàn: Hoá chất dư phải đổ vào bình thu gom chất thải thí nghiệm đúng quy định.'},
            ]
        },
        {
            'context': 'Nung nóng 100\u00A0gam đá vôi (chứa thành phần chính là calcium carbonate, CaCO3) ở nhiệt độ cao, thu được 56\u00A0gam vôi sống (calcium oxide, CaO) và khí carbon dioxide (CO2) bay ra theo phản ứng: CaCO3 → CaO + CO2↑.',
            'items': [
                {'label': 'a', 'statement': 'Quá trình nung đá vôi tạo thành vôi sống và khí CO2 là một hiện tượng biến đổi hoá học.', 'answer': 'ĐÚNG', 'explanation': 'Có sự hình thành các chất mới (CaO và CO2) khác biệt với chất ban đầu (CaCO3).'},
                {'label': 'b', 'statement': 'Theo định luật bảo toàn khối lượng, tổng khối lượng các chất tạo thành luôn lớn hơn khối lượng đá vôi ban đầu.', 'answer': 'SAI', 'explanation': 'Định luật bảo toàn khối lượng khẳng định: Tổng khối lượng sản phẩm luôn bằng tổng khối lượng chất tham gia.'},
                {'label': 'c', 'statement': 'Khối lượng khí carbon dioxide (CO2) thoát ra từ phản ứng nung vôi nói trên là đúng bằng 44\u00A0gam.', 'answer': 'ĐÚNG', 'explanation': 'Bảo toàn khối lượng: m_CO2 = m_CaCO3 - m_CaO = 100 - 56 = 44 gam.'},
                {'label': 'd', 'statement': 'Ở điều kiện chuẩn (25 °C, 1\u00A0bar), thể tích của lượng khí CO2 sinh ra trong phản ứng trên là 24,79\u00A0lít.', 'answer': 'ĐÚNG', 'explanation': 'Số mol CO2: n = 44 / 44 = 1 mol => Thể tích đkc: V = n . 24,79 = 24,79 lít.'},
            ]
        },
        {
            'context': 'Một khối nhôm đặc có thể tích 0,001\u00A0m³ được móc vào một lực kế rồi nhúng chìm hoàn toàn trong nước (cho biết khối lượng riêng của nhôm D_nhôm = 2\u00A0700\u00A0kg/m³, trọng lượng riêng của nước d_nước = 10\u00A0000\u00A0N/m³).',
            'items': [
                {'label': 'a', 'statement': 'Khi khối nhôm chưa nhúng vào nước, số chỉ của lực kế bằng trọng lượng của vật và bằng 27\u00A0N.', 'answer': 'ĐÚNG', 'explanation': 'Khối lượng nhôm: m = D . V = 2 700 . 0,001 = 2,7 kg => Trọng lượng P = 10 . m = 27 N.'},
                {'label': 'b', 'statement': 'Lực đẩy Archimedes của nước tác dụng lên khối nhôm có phương thẳng đứng và chiều từ trên xuống dưới.', 'answer': 'SAI', 'explanation': 'Lực đẩy Archimedes luôn có phương thẳng đứng và chiều từ dưới lên trên.'},
                {'label': 'c', 'statement': 'Độ lớn lực đẩy Archimedes tác dụng lên khối nhôm khi chìm hoàn toàn trong nước là 10\u00A0N.', 'answer': 'ĐÚNG', 'explanation': 'Lực đẩy Archimedes: F_A = d_nước . V = 10 000 . 0,001 = 10 N.'},
                {'label': 'd', 'statement': 'Khi nhúng chìm hoàn toàn trong nước, số chỉ mới của lực kế treo khối nhôm sẽ là 20\u00A0N.', 'answer': 'SAI', 'explanation': 'Số chỉ lực kế khi chìm: F_lk = P - F_A = 27 - 10 = 17 N (không phải 20 N).'},
            ]
        },
    ],
    'essay': [
        {
            'context': 'Để phục vụ tiết thực hành, giáo viên yêu cầu nhóm học sinh pha chế 100\u00A0gam dung dịch muối ăn (sodium chloride, NaCl) nồng độ 10%.',
            'part_a': 'Tính toán khối lượng muối ăn (NaCl) khan và khối lượng nước cất cần dùng để pha chế dung dịch.',
            'part_b': 'Trình bày cụ thể các bước thao tác thực hành pha chế dung dịch trên bằng cân điện tử, ống đong và cốc thủy tinh.',
            'steps': [
                {'content': 'Khối lượng muối ăn (NaCl) cần lấy: m_ct = (m_dd . C%) / 100% = (100 . 10) / 100 = 10 gam.'},
                {'content': 'Khối lượng nước cất cần dùng: m_H2O = m_dd - m_ct = 100 - 10 = 90 gam (tương đương 90 mL nước cất vì D_nước = 1 g/mL).'},
                {'content': 'Thao tác cân: Dùng cân điện tử cân chính xác đúng 10 gam muối ăn NaCl khan cho vào cốc thủy tinh dung tích 250 mL.'},
                {'content': 'Thao tác đong và hòa tan: Dùng ống đong đong đúng 90 mL nước cất rót vào cốc, dùng đũa thủy tinh khuấy đều cho tan hết.'},
            ]
        },
        {
            'context': 'Đốt cháy hoàn toàn 4,8\u00A0gam kim loại magnesium (Mg) trong bình chứa khí oxygen (O2) thu được hợp chất magnesium oxide (MgO) theo phương trình hoá học: 2Mg + O2 → 2MgO.',
            'part_a': 'Tính số mol kim loại Mg đã tham gia phản ứng và tính khối lượng hợp chất MgO tạo thành.',
            'part_b': 'Tính thể tích khí oxygen (O2) tối thiểu cần dùng cho phản ứng trên ở điều kiện chuẩn (25 °C, 1\u00A0bar).',
            'steps': [
                {'content': 'Số mol kim loại Mg: n_Mg = m / M = 4,8 / 24 = 0,2 mol.'},
                {'content': 'Theo PTHH: n_MgO = n_Mg = 0,2 mol. Khối lượng MgO tạo thành: m_MgO = n . M = 0,2 . (24 + 16) = 0,2 . 40 = 8,0 gam.'},
                {'content': 'Theo PTHH: n_O2 = (1/2) . n_Mg = 0,1 mol.'},
                {'content': 'Thể tích khí O2 cần dùng ở đkc: V_O2 = n . 24,79 = 0,1 . 24,79 = 2,479 lít.'},
            ]
        },
        {
            'context': 'Một chiếc xe tải có tổng khối lượng 6\u00A0tấn (6\u00A0000\u00A0kg), được đỡ bởi 6 bánh xe tiếp xúc với mặt đường phẳng. Diện tích tiếp xúc của mỗi bánh xe với mặt đường là 0,025\u00A0m². Lấy gia tốc trọng trường g = 10\u00A0m/s².',
            'part_a': 'Tính trọng lượng của chiếc xe tải và tính tổng diện tích tiếp xúc của cả 6 bánh xe lên mặt đường.',
            'part_b': 'Tính áp suất do xe tải tác dụng lên mặt đường. So sánh với áp suất giới hạn cho phép của mặt đường là 500\u00A0000\u00A0Pa và kết luận.',
            'steps': [
                {'content': 'Trọng lượng xe tải: P = m . g = 6 000 . 10 = 60 000 N (đây chính là áp lực F của xe tác dụng vuông góc lên mặt đường).'},
                {'content': 'Tổng diện tích tiếp xúc của 6 bánh xe: S_tong = 6 . 0,025 = 0,15 m².'},
                {'content': 'Áp suất của xe tải tác dụng lên mặt đường: p = F / S_tong = 60 000 / 0,15 = 400 000 Pa (hoặc 400 000 N/m²).'},
                {'content': 'So sánh: p = 400 000 Pa < p_gioihan = 500 000 Pa. Kết luận: Xe tải được phép lưu thông an toàn mà không làm hỏng mặt đường.'},
            ]
        },
    ]
}


EXAM_DATA_GRADE_9 = {
    'school_info': {
        'dept': 'PHÒNG GD&ĐT HUYỆN LONG PHÚ',
        'school': 'TRƯỜNG THCS LONG PHÚ',
    },
    'exam_info': {
        'title': 'ĐỀ KIỂM TRA ĐỊNH KỲ GIỮA HỌC KỲ I',
        'year': '2026 - 2027',
        'subject': 'KHOA HỌC TỰ NHIÊN',
        'grade': '9',
        'code': '901',
        'duration': 45,
    },
    'mcq': [
        {
            'question': 'Dụng cụ nào sau đây thường được dùng để đo thể tích chất lỏng một cách chính xác trong phòng thực hành môn Khoa học tự nhiên?',
            'options': ['A. Ống đong chia độ.', 'B. Ống nghiệm thủy tinh.', 'C. Đĩa petri nuôi cấy.', 'D. Đèn cồn thí nghiệm.'],
            'answer': 'A'
        },
        {
            'question': 'Dụng cụ nào sau đây thường được dùng để kẹp giữ ống nghiệm khi tiến hành đun nóng hoá chất?',
            'options': ['A. Đũa thủy tinh.', 'B. Kẹp gỗ gắp ống nghiệm.', 'C. Phễu chiết chất lỏng.', 'D. Muỗng lấy hóa chất.'],
            'answer': 'B'
        },
        {
            'question': 'Kí hiệu cảnh báo hình ngọn lửa trên nhãn chai hoá chất mang ý nghĩa cảnh báo điều gì?',
            'options': ['A. Hóa chất ăn mòn da.', 'B. Hóa chất độc sinh học.', 'C. Hóa chất dễ cháy.', 'D. Hóa chất gây nổ mạnh.'],
            'answer': 'C'
        },
        {
            'question': 'Để lấy một lượng nhỏ hoá chất dạng hạt hoặc bột từ lọ đựng, ta nên sử dụng dụng cụ nào?',
            'options': ['A. Ống hút nhỏ giọt.', 'B. Kẹp sắt thí nghiệm.', 'C. Que đóm bằng gỗ.', 'D. Muỗng thủy tinh (thìa).'],
            'answer': 'D'
        },
        {
            'question': 'Khi thực hiện báo cáo khoa học, bước nào sau đây giúp kiểm chứng tính chính xác của giả thuyết đã đặt ra?',
            'options': ['A. Thiết kế bài trình chiếu.', 'B. Thu thập và xử lí số liệu.', 'C. Viết lời cảm ơn và mục lục.', 'D. Chọn màu nền bài báo cáo.'],
            'answer': 'B'
        },
        {
            'question': 'Để bài thuyết trình về một đề tài khoa học tự nhiên đạt hiệu quả cao và thuyết phục người nghe, học sinh nên:',
            'options': ['A. Đọc toàn bộ văn bản trong báo cáo.', 'B. Tránh đưa các kết quả thực nghiệm.', 'C. Kết hợp sơ đồ, hình ảnh trực quan.', 'D. Dùng thật nhiều hiệu ứng chuyển động.'],
            'answer': 'C'
        },
        {
            'question': 'Biểu thức tính động năng của một vật có khối lượng m chuyển động với tốc độ v là',
            'options': ['A. Wđ = (1/2).m.v².', 'B. Wđ = m.v².', 'C. Wđ = (1/2).m².v.', 'D. Wđ = 2.m.v².'],
            'answer': 'A'
        },
        {
            'question': 'Biểu thức tính thế năng trọng trường của vật khối lượng m ở độ cao h so với mốc thế năng là',
            'options': ['A. Wt = (1/2).m.h².', 'B. Wt = m.v.h.', 'C. Wt = P / h.', 'D. Wt = P.h = m.g.h.'],
            'answer': 'D'
        },
        {
            'question': 'Cơ năng của một vật chuyển động trong trọng trường được xác định bằng',
            'options': ['A. Tổng động năng và thế năng.', 'B. Tích giữa động năng và thế năng.', 'C. Hiệu giữa động năng và thế năng.', 'D. Thương giữa thế năng và động năng.'],
            'answer': 'A'
        },
        {
            'question': 'Đơn vị đo động năng, thế năng và cơ năng trong hệ SI là',
            'options': ['A. Oát (W).', 'B. Jun (J).', 'C. Niutơn (N).', 'D. Mã lực (HP).'],
            'answer': 'B'
        },
        {
            'question': 'Một quả bóng rơi tự do từ trên cao xuống đất (bỏ qua lực cản không khí). Trong quá trình rơi:',
            'options': ['A. Động năng giảm, thế năng tăng.', 'B. Cả động năng và thế năng giảm.', 'C. Động năng tăng, thế năng giảm.', 'D. Cả động năng và thế năng tăng.'],
            'answer': 'C'
        },
        {
            'question': 'Khi con lắc đơn dao động từ vị trí cao nhất về vị trí cân bằng, sự chuyển hóa năng lượng chủ yếu là',
            'options': ['A. Động năng chuyển thành thế năng.', 'B. Cơ năng chuyển hoàn toàn thành nhiệt.', 'C. Thế năng chuyển thành hóa năng.', 'D. Thế năng chuyển thành động năng.'],
            'answer': 'D'
        },
        {
            'question': 'Trường hợp nào sau đây lực tác dụng sinh công cơ học?',
            'options': ['A. Lực kéo làm xe lăn bánh trên đường.', 'B. Người đẩy bức tường nhưng tường đứng yên.', 'C. Lực đỡ vật nặng đứng yên trên vai.', 'D. Lực hút Trái Đất tác dụng lên bàn.'],
            'answer': 'A'
        },
        {
            'question': 'Công suất là đại lượng đặc trưng cho',
            'options': ['A. Độ lớn của lực thực hiện công.', 'B. Tốc độ thực hiện công của vật.', 'C. Thời gian kéo dài của chuyển động.', 'D. Quãng đường vật dịch chuyển được.'],
            'answer': 'B'
        },
        {
            'question': 'Một người dùng lực kéo 150\u00A0N để kéo đều một thùng hàng di chuyển quãng đường 10\u00A0m theo hướng của lực. Công thực hiện là',
            'options': ['A. 15\u00A0J.', 'B. 150\u00A0J.', 'C. 1\u00A0500\u00A0J.', 'D. 15\u00A0000\u00A0J.'],
            'answer': 'C'
        },
        {
            'question': 'Một cần cẩu thực hiện công 12\u00A0000\u00A0J để nâng vật nặng lên cao trong thời gian 30\u00A0giây. Công suất của cần cẩu là',
            'options': ['A. 36\u00A0W.', 'B. 120\u00A0W.', 'C. 360\u00A0W.', 'D. 400\u00A0W.'],
            'answer': 'D'
        },
    ],
    'tf': [
        {
            'context': 'Trong tiết thực hành nghiên cứu phản ứng giữa dung dịch hydrochloric acid (HCl) loãng và mẫu đá vôi nhỏ (chứa CaCO3), nhóm học sinh lớp 9 tiến hành chuẩn bị dụng cụ, hoá chất, tiến hành thí nghiệm thu khí CO2 và chuẩn bị bài báo cáo khoa học.',
            'items': [
                {'label': 'a', 'statement': 'Học sinh có thể dùng ống đong chia độ để lấy chính xác thể tích dung dịch HCl cần dùng cho phản ứng.', 'answer': 'ĐÚNG', 'explanation': 'Ống đong chia độ là dụng cụ chuyên dụng đo/lấy thể tích chất lỏng.'},
                {'label': 'b', 'statement': 'Dung dịch acid loãng HCl không độc nên có thể dùng tay cầm trực tiếp vào miệng chai hoá chất để rót.', 'answer': 'SAI', 'explanation': 'Tuyệt đối không dùng tay trần chạm trực tiếp vào hoá chất, phải mang găng tay bảo hộ.'},
                {'label': 'c', 'statement': 'Trong báo cáo khoa học, phần thảo luận kết quả cần đối chiếu thể tích khí CO2 thu được với tính toán lí thuyết.', 'answer': 'ĐÚNG', 'explanation': 'Phần thảo luận cần phân tích, so sánh số liệu thực nghiệm với lí thuyết để rút ra kết luận.'},
                {'label': 'd', 'statement': 'Khi làm bài thuyết trình, chỉ cần trình bày bảng chữ dày đặc kết quả số liệu mà không cần hình ảnh thí nghiệm.', 'answer': 'SAI', 'explanation': 'Bài thuyết trình cần sơ đồ, hình ảnh trực quan, tránh quá nhiều chữ gây nhàm chán.'},
            ]
        },
        {
            'context': 'Một vận động viên trượt ván có khối lượng 50\u00A0kg đang đứng ở đỉnh một dốc trượt cao 3,2\u00A0m so với mặt đất (chọn mốc thế năng tại mặt đất, lấy gia tốc trọng trường g = 10\u00A0m/s², bỏ qua mọi ma sát và lực cản của không khí). Vận động viên bắt đầu trượt không vận tốc đầu xuống chân dốc.',
            'items': [
                {'label': 'a', 'statement': 'Tại vị trí đỉnh dốc, thế năng trọng trường của vận động viên đạt giá trị cực đại và động năng bằng 0.', 'answer': 'ĐÚNG', 'explanation': 'Tại đỉnh dốc (v0 = 0, h max) nên động năng Wđ = 0 và thế năng Wt đạt cực đại.'},
                {'label': 'b', 'statement': 'Tổng động năng và thế năng của vận động viên luôn thay đổi liên tục trong suốt quá trình trượt xuống dốc.', 'answer': 'SAI', 'explanation': 'Bỏ qua ma sát thì cơ năng (tổng động năng và thế năng) được bảo toàn không đổi.'},
                {'label': 'c', 'statement': 'Trong quá trình trượt từ đỉnh dốc xuống mặt đất, thế năng của vận động viên đã chuyển hóa dần thành động năng.', 'answer': 'ĐÚNG', 'explanation': 'Khi trượt xuống, độ cao giảm nên thế năng giảm, tốc độ tăng nên động năng tăng (chuyển hoá).'},
                {'label': 'd', 'statement': 'Tại vị trí chân dốc, toàn bộ thế năng ban đầu chuyển thành động năng, vận tốc đạt được là 8\u00A0m/s.', 'answer': 'ĐÚNG', 'explanation': 'Cơ năng bảo toàn: W = Wt_max = m.g.h = 50.10.3,2 = 1 600 J. Tại chân dốc: (1/2).m.v² = 1600 => v = 8 m/s.'},
            ]
        },
        {
            'context': 'Tại một công trường xây dựng, một động cơ điện được sử dụng để kéo đều một khối vật liệu nặng 200\u00A0kg lên tầng cao có độ cao 12\u00A0m theo phương thẳng đứng trong khoảng thời gian 24\u00A0giây (lấy g = 10\u00A0m/s²).',
            'items': [
                {'label': 'a', 'statement': 'Lực kéo của động cơ tác dụng lên khối vật liệu đã sinh công cơ học vì làm vật dịch chuyển theo phương của lực.', 'answer': 'ĐÚNG', 'explanation': 'Lực kéo tác dụng làm vật dịch chuyển cùng hướng của lực nên đã sinh công cơ học.'},
                {'label': 'b', 'statement': 'Công suất của động cơ càng lớn cho biết động cơ có khả năng thực hiện công càng chậm.', 'answer': 'SAI', 'explanation': 'Công suất là tốc độ thực hiện công, công suất càng lớn thì thực hiện công càng nhanh.'},
                {'label': 'c', 'statement': 'Công cơ học mà động cơ đã thực hiện để đưa khối vật liệu lên độ cao 12\u00A0m là 24\u00A0000\u00A0J.', 'answer': 'ĐÚNG', 'explanation': 'Công nâng đều vật: A = F.s = P.h = m.g.h = 200 . 10 . 12 = 24 000 J.'},
                {'label': 'd', 'statement': 'Công suất nâng vật của động cơ điện trong quá trình nói trên đạt giá trị là 1\u00A0000\u00A0W.', 'answer': 'ĐÚNG', 'explanation': 'Công suất của động cơ: P = A / t = 24 000 / 24 = 1 000 W.'},
            ]
        },
    ],
    'essay': [
        {
            'context': 'Để chuẩn bị cho hội thảo nghiên cứu khoa học cấp trường, nhóm học sinh thực hiện đề tài khảo sát các yếu tố ảnh hưởng đến tốc độ phản ứng hoá học giữa dung dịch hydrochloric acid (HCl) và kim loại kẽm (Zn) tạo thành khí hydrogen (H2).',
            'part_a': 'Em hãy kể tên 02 dụng cụ đo lường cần thiết trong thí nghiệm này và nêu rõ một quy tắc an toàn khi lấy hoá chất acid trong phòng thực hành.',
            'part_b': 'Hãy nêu cấu trúc các bước cơ bản để viết một báo cáo khoa học hoàn chỉnh cho đề tài trên và đề xuất 01 giải pháp giúp bài thuyết trình trước hội đồng trở nên sinh động, lôi cuốn.',
            'steps': [
                {'content': 'Kể tên 02 dụng cụ đo lường phù hợp: Ống đong (đo thể tích dung dịch), cân điện tử (đo khối lượng kẽm) hoặc đồng hồ bấm giây (đo thời gian).'},
                {'content': 'Nêu đúng 01 quy tắc an toàn: Đeo găng tay/kính bảo hộ khi rót acid, không ngửi trực tiếp, đậy kín nắp lọ hoá chất sau khi dùng.'},
                {'content': 'Nêu các bước viết báo cáo khoa học: Tên đề tài -> Mục tiêu -> Dụng cụ, hoá chất -> Tiến trình thí nghiệm -> Kết quả & Thảo luận -> Kết luận.'},
                {'content': 'Đề xuất giải pháp thuyết trình hiệu quả: Sử dụng slide/video ghi lại hiện tượng sủi bọt khí thực tế kết hợp sơ đồ, biểu đồ trực quan.'},
            ]
        },
        {
            'context': 'Một quả bưởi có khối lượng 1,2\u00A0kg rơi tự do từ cành cây ở độ cao 5\u00A0m so với mặt đất (chọn mốc thế năng tại mặt đất, lấy g = 10\u00A0m/s², bỏ qua sức cản của không khí).',
            'part_a': 'Viết biểu thức tính thế năng trọng trường và tính thế năng của quả bưởi tại vị trí cành cây trước khi rơi.',
            'part_b': 'Trong quá trình quả bưởi rơi xuống, có sự chuyển hóa qua lại giữa những dạng năng lượng nào? Cơ năng của quả bưởi khi chạm đất có giá trị bằng bao nhiêu?',
            'steps': [
                {'content': 'Viết đúng biểu thức thế năng trọng trường: Wt = P.h = m.g.h (hoặc nêu rõ các đại lượng kèm đơn vị chuẩn).'},
                {'content': 'Tính đúng thế năng tại cành cây: Wt = 1,2 . 10 . 5 = 60 J.'},
                {'content': 'Nêu đúng sự chuyển hóa năng lượng: Trong quá trình rơi, độ cao giảm nên thế năng chuyển hóa dần thành động năng.'},
                {'content': 'Xác định cơ năng khi chạm đất: Bỏ qua lực cản không khí, cơ năng bảo toàn nên cơ năng chạm đất bằng thế năng ban đầu W = 60 J.'},
            ]
        },
        {
            'context': 'Một thang máy chở người có tổng khối lượng cả cabin và hành khách là 600\u00A0kg, được kéo chuyển động đều đi lên tầng cao quãng đường 15\u00A0m trong thời gian 20\u00A0giây theo phương thẳng đứng (lấy g = 10\u00A0m/s²).',
            'part_a': 'Nêu định nghĩa công suất và giải thích ý nghĩa vật lí của công suất trong trường hợp chuyển động của thang máy trên.',
            'part_b': 'Tính công cơ học do lực kéo của động cơ thang máy thực hiện và tính công suất của động cơ trong quá trình nâng thang máy.',
            'steps': [
                {'content': 'Nêu đúng định nghĩa công suất: Công suất là đại lượng đặc trưng cho tốc độ thực hiện công, được xác định bằng công thực hiện trong một đơn vị thời gian (P = A / t).'},
                {'content': 'Giải thích ý nghĩa: Cho biết mức độ hoạt động mạnh/yếu của động cơ thang máy (trong 1 giây thang máy thực hiện được bao nhiêu Jun công).'},
                {'content': 'Tính công cơ học của động cơ: Trọng lượng P = m.g = 600 . 10 = 6 000 N. Công nâng đều: A = F.s = 6 000 . 15 = 90 000 J.'},
                {'content': 'Tính công suất của động cơ thang máy: P = A / t = 90 000 / 20 = 4 500 W (hoặc 4,5 kW).'},
            ]
        },
    ]
}


def generate_exam(exam_data, output_filepath):
    doc = create_document()
    add_exam_header(doc, exam_data['school_info'], exam_data['exam_info'])
    add_part1_mcq(doc, exam_data['mcq'])
    add_part2_true_false(doc, exam_data['tf'])
    add_part3_essay(doc, exam_data['essay'])

    mcq_answers = [q['answer'] for q in exam_data['mcq']]
    add_part2_answers_and_rubric(doc, mcq_answers, exam_data['tf'], exam_data['essay'])

    doc.save(output_filepath)
    print(f"✓ Đã lưu thành công file: {output_filepath}")


def verify_generated_docx(filepath):
    doc = docx.Document(filepath)
    sec = doc.sections[0]
    w_cm = round(sec.page_width.cm, 1)
    h_cm = round(sec.page_height.cm, 1)
    left_cm = round(sec.left_margin.cm, 1)
    right_cm = round(sec.right_margin.cm, 1)
    top_cm = round(sec.top_margin.cm, 1)
    bottom_cm = round(sec.bottom_margin.cm, 1)

    print(f"\n--- KIỂM TRA FILE: {os.path.basename(filepath)} ---")
    print(f"• Khổ giấy: {w_cm} x {h_cm} cm (A4 chuẩn: 21.0 x 29.7 cm)")
    print(f"• Lề: Trái {left_cm} cm, Phải {right_cm} cm, Trên {top_cm} cm, Dưới {bottom_cm} cm")
    print(f"• Số đoạn văn: {len(doc.paragraphs)}, Số bảng: {len(doc.tables)}")

    non_12_runs = 0
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.size and r.font.size != Pt(12):
                non_12_runs += 1
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.font.size and r.font.size != Pt(12):
                            non_12_runs += 1

    print(f"• Kiểm tra cỡ chữ 12 pt: {'ĐẠT (100% runs)' if non_12_runs == 0 else f'Cảnh báo: {non_12_runs} runs khác 12pt'}")


def main():
    exams = [
        (EXAM_DATA_GRADE_7, "De_KT_GK1_KHTN_Lop7_Ma701.docx"),
        (EXAM_DATA_GRADE_8, "De_KT_GK1_KHTN_Lop8_Ma801.docx"),
        (EXAM_DATA_GRADE_9, "De_KT_GK1_KHTN_Lop9_Ma901.docx"),
    ]

    for data, fname in exams:
        ans = [q['answer'] for q in data['mcq']]
        cA = ans.count('A')
        cB = ans.count('B')
        cC = ans.count('C')
        cD = ans.count('D')
        grade = data['exam_info']['grade']
        print(f"Khối {grade} - Phân bố khóa trắc nghiệm: A={cA}, B={cB}, C={cC}, D={cD} (Tổng: {len(ans)})")
        assert len(ans) == 16, f"Khối {grade}: Số câu không phải 16"
        assert cA == 4 and cB == 4 and cC == 4 and cD == 4, f"Khối {grade}: Khóa không cân bằng 4A-4B-4C-4D"

        generate_exam(data, fname)
        verify_generated_docx(fname)

        public_dir = os.path.join(os.getcwd(), "public")
        if os.path.exists(public_dir):
            import shutil
            shutil.copy(fname, os.path.join(public_dir, fname))
            print(f"  -> Đã sao chép vào public/{fname}")

    print("\n🎉 HOÀN TẤT TẠO 3 FILE WORD ĐỀ KIỂM TRA ĐỊNH KỲ GIỮA HỌC KỲ I!")


if __name__ == "__main__":
    main()
